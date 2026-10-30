"""Create non-personal DOCX sources for native Word clipboard acceptance."""
import json
from pathlib import Path
import sys
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

directory = Path(sys.argv[1]); directory.mkdir(parents=True, exist_ok=True)
def base(title):
    doc = Document(); section = doc.sections[0]; section.page_width = Mm(210); section.page_height = Mm(297)
    section.top_margin = section.bottom_margin = Mm(20); section.left_margin = section.right_margin = Mm(20)
    style = doc.styles['Normal']; style.font.name = 'Arial'; style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    for name in ['Title', 'Heading 1', 'Heading 2']:
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.core_properties.author = 'ETranslate QA'; doc.core_properties.last_modified_by = 'ETranslate QA'; doc.core_properties.title = title
    doc.add_paragraph(title, 'Title')
    return doc
def rtl(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    bidi = OxmlElement('w:bidi'); p._p.get_or_add_pPr().append(bidi)
def paint(cell, fill):
    props = cell._tc.get_or_add_tcPr(); shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), fill); props.append(shade)
    borders = OxmlElement('w:tcBorders')
    for side in ['top', 'left', 'bottom', 'right']:
        edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '8'); edge.set(qn('w:color'), 'D9D9D9'); borders.append(edge)
    props.append(borders)

doc = base('Word clipboard typography test')
doc.add_paragraph('RW_TYPO_START İstanbul Iı İi Şş Ğğ Çç Öö Üü $5 <sample>')
p = doc.add_paragraph(); rtl(p); p.paragraph_format.line_spacing = 1.5; p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(12); p.paragraph_format.first_line_indent = Mm(5)
r = p.add_run('RW_STYLE ترجمهٔ آزمایشی با نیم‌فاصله İstanbul'); r.font.name = 'Times New Roman'; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = RGBColor.from_string('B91C1C'); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
p = doc.add_paragraph('RW_INLINE ')
for text, name in [('italic ', 'italic'), ('underline ', 'underline'), ('strike ', 'strike'), ('SUP ', 'superscript'), ('SUB', 'subscript')]:
    run = p.add_run(text); setattr(run.font, name, True)
doc.add_paragraph('RW_EMPTY_BEFORE'); doc.add_paragraph(); doc.add_paragraph('RW_EMPTY_AFTER')
p = doc.add_paragraph('RW_PAGE_TWO'); p.paragraph_format.page_break_before = True
doc.add_paragraph('این متن آزمایشی است و هیچ اطلاعات شخصی ندارد. ترجمهٔ فارسی با نیم‌فاصله.')
doc.add_paragraph('RW_TYPO_END Bu belge yalnızca çeviri düzenleyicisini sınamak içindir.')
doc.save(directory / 'typography.docx')

doc = base('Word clipboard tables and lists test')
doc.add_paragraph('RW_TABLE_START')
for i in range(1, 4): doc.add_paragraph(f'RW_NUMBER_{i:02} Türkçe liste öğesi', 'List Number')
for i in range(1, 3): doc.add_paragraph(f'RW_BULLET_{i:02} فهرست آزمایشی', 'List Bullet')
table = doc.add_table(rows=1, cols=2); table.autofit = False; table.columns[0].width = Mm(100); table.columns[1].width = Mm(65)
for i, text in enumerate(['RW_HEADER_A Açıklama', 'RW_HEADER_B Değer']): table.rows[0].cells[i].text = text; paint(table.rows[0].cells[i], 'DBEAFE')
repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
for i in range(1, 27):
    cells = table.add_row().cells
    cells[0].text = f'RW_ROW_{i:02} İstanbul çeviri metni ve tablo satırı. ' + 'Uzun açıklama ve karşılaştırma. ' * 2
    cells[1].text = f'RW_VALUE_{i:02} نمونهٔ ترجمه'
    if i == 4: cells[0].merge(cells[1])
    for cell in table.rows[-1].cells: paint(cell, 'FFFFFF' if i % 2 else 'F3F4F6')
doc.add_paragraph('RW_TABLE_END')
doc.save(directory / 'tables.docx')

doc = base('Word clipboard unsupported features test'); doc.add_paragraph('RW_UNSUPPORTED_START')
image = directory / 'sample-image.png'; Image.new('RGB', (180, 90), '#1d4ed8').save(image)
run = doc.add_paragraph('RW_IMAGE_BEFORE ').add_run(); drawing = run.add_picture(str(image), width=Mm(35)); drawing._inline.docPr.set('descr', 'RW_IMAGE_ALT')
doc.add_paragraph('RW_IMAGE_AFTER')
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = 'RW_VERTICAL_A'; table.cell(1, 0).text = 'RW_VERTICAL_B'; table.cell(0, 0).merge(table.cell(1, 0))
table.cell(0, 1).text = 'RW_RIGHT_A'; table.cell(1, 1).text = 'RW_RIGHT_B'
p = doc.add_paragraph('RW_LINK '); link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), doc.part.relate_to('https://example.invalid/word-qa', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
r = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'RW_LINK_TEXT'; r.append(text); link.append(r); p._p.append(link)
doc.add_paragraph('RW_UNSUPPORTED_END'); doc.save(directory / 'unsupported.docx')
print(json.dumps({'sources': ['typography.docx', 'tables.docx', 'unsupported.docx'], 'personalData': False}))
